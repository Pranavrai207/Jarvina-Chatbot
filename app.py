import os
import json
import re
import logging
from contextlib import asynccontextmanager
from fastapi import FastAPI, Request, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict
from typing import Optional
from datetime import datetime
import io

# --- New Imports for File Generation ---
from fpdf import FPDF
from openpyxl import Workbook
from docx import Document

# Import the new GeminiClient from the updated llm_client.py
from llm_client import GeminiClient

# --- Import the Database class ---
from database import Database

# Load environment variables from .env file
from dotenv import load_dotenv
load_dotenv()

# Initialize globals
gemini_client: Optional[GeminiClient] = None
jarvina_persona_data: Optional[dict] = None
custom_replies_map: dict = {}
CUSTOM_INSTRUCTIONS_FILE = "custom_instructions.json"
db: Optional[Database] = None 

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

# --- Helper Functions ---
def normalize_text(text: str) -> str:
    """
    Converts text to lowercase, removes punctuation, and normalizes spaces.
    """
    text = text.lower()
    text = re.sub(r'[^\w\s]', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def load_custom_instructions_from_file() -> str:
    """Loads custom instructions from a JSON file."""
    if os.path.exists(CUSTOM_INSTRUCTIONS_FILE):
        try:
            with open(CUSTOM_INSTRUCTIONS_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return data.get("instructions", "")
        except json.JSONDecodeError as e:
            logging.error(f"Error decoding {CUSTOM_INSTRUCTIONS_FILE}: {e}")
            return ""
        except Exception as e:
            logging.error(f"Error loading {CUSTOM_INSTRUCTIONS_FILE}: {e}")
            return ""
    return ""

def save_custom_instructions_to_file(instructions: str):
    """Saves custom instructions to a JSON file."""
    try:
        with open(CUSTOM_INSTRUCTIONS_FILE, 'w', encoding='utf-8') as f:
            json.dump({"instructions": instructions}, f, indent=4)
    except Exception as e:
        logging.error(f"Error saving to {CUSTOM_INSTRUCTIONS_FILE}: {e}")
        raise

# --- Lifespan Event Handler ---
@asynccontextmanager
async def lifespan(app: FastAPI):
    global gemini_client
    global jarvina_persona_data
    global custom_replies_map
    global db

    # 1. Initialize Gemini Client
    try:
        gemini_client = GeminiClient()
        print("GeminiClient initialized successfully.")
    except ValueError as e:
        print(f"Error during client initialization: {e}")
        gemini_client = None

    # 2. Initialize Database
    db = Database()
    print("Database connection established.")

    # 3. Load Persona & Memory
    memory_file_path = "memory.json"
    if os.path.exists(memory_file_path):
        try:
            with open(memory_file_path, 'r', encoding='utf-8') as f:
                memory_data = json.load(f)

            jarvina_persona_data = memory_data.get("persona_data")
            if jarvina_persona_data:
                print(f"Persona data loaded from {memory_file_path} successfully.")
            else:
                print(f"Warning: 'persona_data' not found in {memory_file_path}. Using default persona.")

            loaded_custom_replies = memory_data.get("custom_replies", [])
            for reply_entry in loaded_custom_replies:
                response = reply_entry.get("response")
                phrases = reply_entry.get("phrases", [])
                if response and phrases:
                    for phrase in phrases:
                        custom_replies_map[normalize_text(phrase)] = response
            
            if custom_replies_map:
                print(f"Custom replies loaded from {memory_file_path} successfully.")
            else:
                print(f"Warning: 'custom_replies' not found or empty in {memory_file_path}.")

        except json.JSONDecodeError as e:
            print(f"Error decoding memory.json: {e}. Persona data and custom replies will not be used.")
            jarvina_persona_data = None
            custom_replies_map = {}
        except Exception as e:
            print(f"Error loading memory.json: {e}. Persona data and custom replies will not be used.")
            jarvina_persona_data = None
            custom_replies_map = {}
    else:
        print(f"Warning: {memory_file_path} not found. Using default persona.")
        jarvina_persona_data = None
        custom_replies_map = {}

    yield
    print("FastAPI application shutting down.")

# --- App Initialization ---
app = FastAPI(lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

templates = Jinja2Templates(directory="templates")

# --- Pydantic Models ---
class ChatRequest(BaseModel):
    messages: Optional[list[dict]] = None
    prompt: Optional[str] = None
    model_config = ConfigDict(extra='allow')

class CustomInstructionsRequest(BaseModel):
    instructions: str

class TextRequest(BaseModel):
    text: str

# --- File Generation Endpoints ---
@app.post("/generate-pdf")
async def generate_pdf(request: TextRequest):
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, request.text)
        # Output to string/bytes
        pdf_output = pdf.output(dest='S').encode('latin1', errors='replace')
        buffer = io.BytesIO(pdf_output)
        return StreamingResponse(
            buffer, 
            media_type='application/pdf', 
            headers={'Content-Disposition': 'attachment; filename=output.pdf'}
        )
    except Exception as e:
        logging.error(f"PDF Generation Error: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to generate PDF: {e}")

@app.post("/generate-excel")
async def generate_excel(request: TextRequest):
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Generated Text"
        for row_index, line in enumerate(request.text.splitlines(), 1):
            sheet.cell(row=row_index, column=1, value=line)
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer, 
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 
            headers={'Content-Disposition': 'attachment; filename=output.xlsx'}
        )
    except Exception as e:
        logging.error(f"Excel Generation Error: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel: {e}")

@app.post("/generate-word")
async def generate_word(request: TextRequest):
    try:
        document = Document()
        document.add_paragraph(request.text)
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer, 
            media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', 
            headers={'Content-Disposition': 'attachment; filename=output.docx'}
        )
    except Exception as e:
        logging.error(f"Word Generation Error: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to generate Word doc: {e}")

# --- General Chat Endpoint ---
@app.post("/api/chat/mistral/")
async def chat_endpoint(chat_request: ChatRequest):
    if not gemini_client:
        raise HTTPException(status_code=500, detail="GeminiClient not initialized.")

    user_prompt = chat_request.prompt or ""
    
    # Save user message to DB if it exists
    if user_prompt:
        db.add_conversation('user', user_prompt)

    # --- Command Handling ---
    normalized_prompt = user_prompt.lower().strip()

    if normalized_prompt == "clear history":
        db.clear_history_database()
        response_text = "Chat history has been cleared."
        db.add_conversation('assistant', response_text)
        return JSONResponse(content={"response": response_text})
    
    if normalized_prompt == "clear notes":
        db.clear_notes_database()
        response_text = "All saved notes have been cleared."
        db.add_conversation('assistant', response_text)
        return JSONResponse(content={"response": response_text})

    if normalized_prompt.startswith("save:"):
        note_content = user_prompt[len("save:"):].strip()
        if note_content:
            db.add_note(note_content)
            response_text = f"Got it. I've saved the note: \"{note_content}\""
            db.add_conversation('assistant', response_text)
            return JSONResponse(content={"response": response_text})
        else:
            response_text = "Please provide content after 'save:' to save a note."
            db.add_conversation('assistant', response_text)
            return JSONResponse(content={"response": response_text})

    # --- Custom & Pre-programmed Replies ---
    if "current time" in normalized_prompt or "what time is it" in normalized_prompt:
        current_time_str = datetime.now().strftime('%I:%M:%S %p on %A, %B %d, %Y')
        response_text = f"The current time is {current_time_str}."
        db.add_conversation('assistant', response_text)
        return JSONResponse(content={"response": response_text})

    normalized_for_reply_map = normalize_text(user_prompt)
    if normalized_for_reply_map in custom_replies_map:
        response_text = custom_replies_map[normalized_for_reply_map]
        db.add_conversation('assistant', response_text)
        return JSONResponse(content={"response": response_text})

    # --- LLM Generation ---
    try:
        # 1. Build the System Prompt
        system_content_parts = ["CRITICAL INSTRUCTION: DO NOT use emojis."]
        custom_instructions = load_custom_instructions_from_file()
        if custom_instructions:
            system_content_parts.append(f"User's Custom Instructions: {custom_instructions}")
        
        if jarvina_persona_data:
             system_content_parts.append(json.dumps(jarvina_persona_data))
        else:
            system_content_parts.append("You are Jarvina, a personal AI assistant.")
        
        system_message_text = "\n".join(system_content_parts)

        # 2. Build the Message List for Gemini
        gemini_formatted_messages = []

        # Inject System Prompt as a fake User/Model exchange
        gemini_formatted_messages.append({
            "role": "user", 
            "parts": [{"text": f"SYSTEM INSTRUCTIONS:\n{system_message_text}"}]
        })
        gemini_formatted_messages.append({
            "role": "model", 
            "parts": [{"text": "Understood. I will follow these instructions."}]
        })

        # 3. Append Conversation History
        if chat_request.messages:
            for message in chat_request.messages:
                role = "model" if message.get("role") == "assistant" else "user"
                content = message.get("content")
                if content:
                    gemini_formatted_messages.append({"role": role, "parts": [{"text": content}]})
        
        # 4. Append the Latest User Prompt (if not already last)
        should_append_prompt = True
        if gemini_formatted_messages:
            last_message_text = gemini_formatted_messages[-1]["parts"][0]["text"]
            if last_message_text == user_prompt:
                should_append_prompt = False
        
        if user_prompt and should_append_prompt:
             gemini_formatted_messages.append({"role": "user", "parts": [{"text": user_prompt}]})

        # 5. Call the Client
        response_content = gemini_client.generate_content(messages=gemini_formatted_messages)
        
        # Save Assistant Response to DB
        db.add_conversation('assistant', response_content)

        return JSONResponse(content={"response": response_content})
    except Exception as e:
        logging.error(f"Error during Gemini generation: {e}")
        raise HTTPException(status_code=500, detail=str(e))

# --- Other Endpoints ---
@app.post("/api/save_custom_instructions")
async def save_custom_instructions(request: CustomInstructionsRequest):
    try:
        save_custom_instructions_to_file(request.instructions)
        return JSONResponse(content={"message": "Custom instructions saved successfully!"})
    except Exception as e:
        raise HTTPException(status_code=500, detail="Failed to save custom instructions.")

@app.get("/api/load_custom_instructions")
async def load_custom_instructions_api():
    instructions = load_custom_instructions_from_file()
    return JSONResponse(content={"instructions": instructions})

@app.get("/", response_class=HTMLResponse)
async def read_root(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})
